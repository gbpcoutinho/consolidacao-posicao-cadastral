# Formatacao e exportacao do Excel (secao 8 do TRATAMENTO.R).
import io
import math
import os
import tempfile
import warnings
from copy import copy
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .config import CODIGO_UNICO, LARGURA_MANUAL
from .erros import ErroValidacao
from .protecao import cifrar_para

# O Excel aceita ate 255; textos longos ficam legiveis com quebra a partir daqui.
LARGURA_AUTO_MAXIMA = 100

# O Excel tem 1.048.576 linhas por aba; a primeira e o cabecalho.
MAXIMO_LINHAS_DADOS = 1_048_575

ABAS_SAIDA = [
    ("consolidado", "Base Tratada", "TabelaBaseTratada"),
    ("pendencias", "Pendencias", "TabelaPendencias"),
    ("log", "Log", "TabelaLog"),
]

COLUNAS_CENTRALIZADAS = [
    "EMPRESA", "SubF", "NUMERO DO CERTIFICADO", CODIGO_UNICO,
    "TITULARIDADE", "NUMERO DO CARTAO", "PLANO", "REGIAO",
    "GRAU PARENTESCO",
    "SEXO", "ESTADO CIVIL", "NUMERO DO CPF", "CARGO DE OCUPACAO",
    "NUMERO DA MATRICULA",
]

_ESTILO_CABECALHO = {
    "font": Font(color="FFFFFF", bold=True),
    "fill": PatternFill("solid", fgColor="4472C4"),
    "alignment": Alignment(horizontal="center", vertical="center", wrap_text=True),
    "border": Border(bottom=Side(style="thin")),
}
_ALINHAMENTO_CORPO = Alignment(vertical="center")
_ALINHAMENTO_CENTRO = Alignment(horizontal="center", vertical="center")
_ALINHAMENTO_DIREITA = Alignment(horizontal="right", vertical="center")


def _texto_exibido(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (date, datetime)):
        return "00/00/0000"
    return str(v)


def _linhas_quebradas(texto: str, largura: int) -> int:
    """Quantas linhas o texto ocupa com quebra por palavra (palavra longa e cortada)."""
    linhas, atual = 0, 0
    for palavra in texto.split():
        while len(palavra) > largura:
            if atual:
                linhas, atual = linhas + 1, 0
            linhas += 1
            palavra = palavra[largura:]
        if atual and atual + 1 + len(palavra) <= largura:
            atual += 1 + len(palavra)
        else:
            if atual:
                linhas += 1
            atual = len(palavra)
    return linhas + (1 if atual else 0)


def largura_coluna_data(nome: str) -> float:
    # Pedido do usuario (diferente do R, que fixa 13): a coluna de data fica
    # larga o bastante para o cabecalho caber em ate 2 linhas.
    caracteres = 1
    while _linhas_quebradas(nome, caracteres) > 2:
        caracteres += 1
    # Maiusculas em negrito ocupam ~25% a mais que a unidade de largura do Excel.
    return max(13, math.ceil(caracteres * 1.25 + 2))


def _coluna_de_datas(nome: str, valores: list) -> bool:
    # No R, o estilo de data vale para colunas com DATA no nome e classe Date,
    # mesmo vazias. Vazia sem tipo: conta como data se o nome comeca com DATA.
    if "DATA" not in nome.upper() or not all(v is None or isinstance(v, date) for v in valores):
        return False
    return any(isinstance(v, date) for v in valores) or nome.upper().startswith("DATA")


def _celula(v):
    # openpyxl recusa caracteres de controle; o Excel tambem nao os aceita.
    return ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v


def _alinhamento_coluna(nome: str, data: bool) -> Alignment:
    if data:
        return _ALINHAMENTO_CENTRO
    if nome == "CODIGO DO DEPENDENTE":
        return _ALINHAMENTO_DIREITA
    if nome in COLUNAS_CENTRALIZADAS:
        return _ALINHAMENTO_CENTRO
    return _ALINHAMENTO_CORPO


def adicionar_aba_formatada(wb, df: pd.DataFrame, nome_aba: str, nome_tabela: str) -> None:
    """Escreve a aba num Workbook(write_only=True): estilos, larguras, altura do
    cabecalho, tabela/filtro e congelamento sao definidos antes das linhas."""
    if len(df.columns) == 0:
        raise ValueError(f"A aba '{nome_aba}' nao tem colunas para exportar.")
    ws = wb.create_sheet(nome_aba)
    ws.sheet_view.showGridLines = False

    colunas = [str(c) for c in df.columns]
    n_linhas, n_colunas = len(df), len(colunas)
    valores = [df[c].tolist() for c in df.columns]

    for j, nome in enumerate(colunas, start=1):
        largura = min(
            LARGURA_AUTO_MAXIMA,
            max(len(_texto_exibido(v)) for v in [nome, *valores[j - 1]]) + 2,
        )
        if nome in LARGURA_MANUAL:
            largura = LARGURA_MANUAL[nome]
        # A mascara dd/mm/yyyy precisa de largura minima para nao aparecer como ####.
        if "DATA" in nome.upper():
            largura = largura_coluna_data(nome)
        ws.column_dimensions[get_column_letter(j)].width = largura
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"

    ultima = get_column_letter(n_colunas)
    if n_linhas > 0:
        tabela = Table(displayName=nome_tabela, ref=f"A1:{ultima}{n_linhas + 1}")
        tabela.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        # No modo write_only o openpyxl nao le os cabecalhos da aba: nomes explicitos.
        tabela._initialise_columns()
        for coluna, nome in zip(tabela.tableColumns, colunas):
            coluna.name = nome
        with warnings.catch_warnings():
            # Aviso fixo do modo write_only; as colunas ja foram nomeadas acima.
            warnings.filterwarnings("ignore", message="In write-only mode")
            ws.add_table(tabela)
    else:
        ws.auto_filter.ref = f"A1:{ultima}1"

    cabecalho = []
    for nome in colunas:
        celula = WriteOnlyCell(ws, value=nome)
        for atributo, estilo in _ESTILO_CABECALHO.items():
            setattr(celula, atributo, estilo)
        cabecalho.append(celula)
    ws.append(cabecalho)

    # Um molde de estilo por coluna: atribuir alinhamento celula a celula
    # recalcula o estilo a cada vez e dominava o tempo de gravacao.
    moldes = []
    for nome, coluna in zip(colunas, valores):
        data = _coluna_de_datas(nome, coluna)
        molde = WriteOnlyCell(ws)
        molde.alignment = _alinhamento_coluna(nome, data)
        if data:
            molde.number_format = "dd/mm/yyyy"
        moldes.append(molde._style)

    for i in range(n_linhas):
        linha = []
        for j in range(n_colunas):
            celula = WriteOnlyCell(ws, value=_celula(valores[j][i]))
            celula._style = copy(moldes[j])
            linha.append(celula)
        ws.append(linha)


def validar_limite_linhas(c) -> None:
    """Acima do limite o Excel abre o arquivo pedindo reparo e perde linhas."""
    for atributo, nome_aba, _ in ABAS_SAIDA:
        n = len(getattr(c, atributo))
        if n > MAXIMO_LINHAS_DADOS:
            raise ErroValidacao(
                f"A aba '{nome_aba}' teria {n:,} linhas de dados; o Excel aceita no maximo "
                f"{MAXIMO_LINHAS_DADOS:,}. Processe os arquivos em lotes menores.".replace(",", ".")
            )


def escrever_excel(c, destino: Path, senha: str | None = None) -> None:
    """Grava o consolidado em `destino`. Com `senha`, o arquivo e montado em
    memoria e so a versao cifrada chega ao disco (vazia ou None = sem senha)."""
    validar_limite_linhas(c)
    # write_only: grava em fluxo, sem montar a planilha inteira na memoria.
    wb = openpyxl.Workbook(write_only=True)
    for atributo, nome_aba, nome_tabela in ABAS_SAIDA:
        adicionar_aba_formatada(wb, getattr(c, atributo), nome_aba, nome_tabela)

    # Grava num temporario na mesma pasta e so entao troca: se o destino estiver
    # aberto no Excel, o arquivo anterior fica intacto e o erro e claro.
    # Prefixo "~$": se o temporario sobrar (processo morto), a listagem de
    # entradas do R e do Python ja o ignora, como os arquivos de bloqueio do Excel.
    destino = Path(destino)
    temporario = None
    try:
        destino.parent.mkdir(parents=True, exist_ok=True)
        descritor, nome = tempfile.mkstemp(dir=destino.parent, prefix="~$", suffix=".xlsx")
        os.close(descritor)
        temporario = Path(nome)
        if senha:
            plano = io.BytesIO()
            wb.save(plano)
            cifrar_para(plano, senha, temporario)
        else:
            wb.save(temporario)
        os.replace(temporario, destino)
        temporario = None
    except OSError as e:
        raise ErroValidacao(
            f"Nao foi possivel salvar o arquivo de saida '{destino}': feche o arquivo "
            f"se estiver aberto no Excel e tente novamente. ({e})"
        ) from e
    finally:
        # Qualquer falha (inclusive fora de OSError) nao deixa o temporario na pasta.
        if temporario is not None:
            _descartar_abas(wb)
            try:
                temporario.unlink(missing_ok=True)
            except OSError:
                pass  # nao esconde o erro original


def _descartar_abas(wb) -> None:
    # Abas write_only que nao chegaram a ser gravadas deixam um gerador aberto;
    # fecha-lo aqui evita um aviso "Exception ignored" mais tarde.
    for ws in wb.worksheets:
        linhas = getattr(ws, "_rows", None)
        if linhas is not None:
            try:
                linhas.close()
            except Exception:
                pass
