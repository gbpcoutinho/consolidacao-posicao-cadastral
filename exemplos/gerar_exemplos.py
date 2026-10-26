"""Gera 3 planilhas de posicao cadastral FICTICIAS no formato do arquivo real.

Modelo: "PC73301 10.2026 - sem dados pessoais.xlsx" (5 abas, tudo como texto,
datas dd/mm/aa + colunas (Y2K), "00/00/0000" para data vazia, registros
tipo 1..9 e as 2 linhas finais de totais na aba de dependentes).

Uso: .venv\\Scripts\\python exemplos\\gerar_exemplos.py
Saida: exemplos/entrada/PC9000N 10.2026.xlsx e exemplos/COD_EMPRESAS_EXEMPLO.csv
Nenhum dado real: nomes, CPFs, cartoes, contas e empresas sao inventados.
"""
import random
from datetime import date, timedelta
from pathlib import Path

import openpyxl

PASTA = Path(__file__).resolve().parent
COMPETENCIA = date(2026, 9, 8)
DATA_VAZIA = "00/00/0000"

EMPRESAS = [
    ("90001", "ALFA LOGISTICA FICTICIA LTDA", 25),
    ("90002", "BETA SERVICOS EXEMPLO S A", 15),
    ("90003", "GAMA INDUSTRIA DEMONSTRACAO", 10),
]

NOMES_F = ["ANA", "BEATRIZ", "CARLA", "DANIELA", "ELISA", "FERNANDA", "GABRIELA", "HELENA", "ISABEL", "JULIA"]
NOMES_M = ["ANDRE", "BRUNO", "CARLOS", "DIEGO", "EDUARDO", "FELIPE", "GUSTAVO", "HENRIQUE", "IGOR", "JOAO"]
SOBRENOMES = ["SILVA", "SOUZA", "OLIVEIRA", "PEREIRA", "COSTA", "RODRIGUES", "ALMEIDA", "NASCIMENTO", "LIMA", "ARAUJO"]
CARGOS = ["ANALISTA ADMINISTRATIVO PL", "ASSISTENTE DE VENDAS", "COORDENADOR DE QSMS",
          "GERENTE DE VENDAS SR", "TECNICO DE SEGURANCA", "OPERADOR DE LOGISTICA"]
PLANOS = ["TNQ2", "NP06", "TNQ1"]
BANCOS = [("341", "9674"), ("001", "3021"), ("237", "1520")]

CAB_TITULAR_ARQUIVO = [
    "TIPO DE REGISTRO", "TIPO DE ARQUIVO", "NÚMERO DO CONTRATO", "DATA DE COMPETÊNCIA",
    "DATA DE PROCESSAMENTO", "DATA DE COMPETÊNCIA (Y2K)", "DATA DE PROCESSAMENTO (Y2K)", "FILLER",
]
CAB_TITULAR = [
    "TIPO DE REGISTRO", "NUMERO DA SUBFATURA", "NUMERO DO CERTIFICADO", "NOME DO SEGURADO",
    "NUMERO DA MATRICULA", "SEXO DO SEGURADO", "DATA DE NASCIMENTO", "ESTADO CIVIL",
    "NUMERO DO CPF", "CARGO DE OCUPAÇÃO", "DATA DE ADMISSAO", "DATA DE INICIO DE VIGENCIA",
    "PLANO", "MATRICULA ESPECIAL", "DATA DE NASCIMENTO (Y2K)", "DATA DE ADMISSAO (Y2K)",
    "DATA DE INICIO DE VIGENCIA (Y2K)", "AREA RESERVADA", "DATA DE REATIVAÇÃO", "REGIÃO",
    "DATA DE CANCELAMENTO (Y2K)", "FILLER",
]
CAB_PIS = [
    "TIPO DE REGISTRO", "NUMERO DA SUBFATURA", "NUMERO DO CERTIFICADO", "NOME DO SEGURADO",
    "NUMERO DA MATRICULA", "NATUREZA DOC. IDENTIDADE", "NUMERO DOC. IDENTIDADE",
    "NOME ÓRGÃO EXPEDITOR", "CODIGO PAIS", "DATA EXPEDIÇÃO IDENTIDADE", "NUMERO DO PIS", "FILLER",
]
CAB_NOME_MAE = [
    "TIPO DE REGISTRO", "NUMERO DA SUBFATURA", "NUMERO DO CERTIFICADO", "NOME DO SEGURADO",
    "NUMERO DA MATRICULA", "CARTÃO", "NOME DA MÃE", "MATRICULA ESPECIAL", "FILLER",
]
CAB_BANCO = [
    "TIPO DE REGISTRO", "NUMERO DA SUBFATURA", "NUMERO DO CERTIFICADO", "NOME DO SEGURADO",
    "NUMERO DA MATRICULA", "BANCO", "AGENCIA", "PRIMEIRO DV DA AGENCIA", "SEGUNDO DV DA AGENCIA",
    "CONTA CORRENTE", "PRIMEIRO DV DA CONTA", "SEGUNDO DV DA CONTA", "MATRICULA ESPECIAL", "FILLER",
]
CAB_DEPENDENTES = [
    "TIPO DE REGISTRO", "NUMERO DA SUBFATURA", "NUMERO DO CERTIFICADO", "CODIGO DO DEPENDENTE",
    "NOME DEPENDENTE", "DATA DE NASCIMENTO", "SEXO DO DEPENDENTE", "ESTADO CIVIL DO DEPENDENTE",
    "GRAU PARENTESCO", "DATA DE INICIO DE VIGENCIA", "MATRICULA ESPECIAL",
    "DATA DE NASCIMENTO (Y2K)", "DATA DE INICIO DE VIGENCIA (Y2K)",
    "MATRICULA ESPECIAL DO DEPENDENTE", "DATA DE REATIVAÇÃO", "DATA DE CANCELAMENTO",
    "NUMERO DO CPF", "NUMERO DO CARTÃO DEPENDENTE", "NOME DA MÃE DO DEPENDENTE",
]
CAB_TOTAIS = [
    "TIPO DE REGISTRO", "TOTAL DE TITULARES", "TOTAL DE DADOS BANCÁRIOS", "TOTAL DE CORRENTISTAS",
    "TOTAL DE DEPENDENTES", "TOTAL DE REGISTROS", "FILLER",
]


def _nome(rnd, sexo):
    primeiro = rnd.choice(NOMES_F if sexo == "02" else NOMES_M)
    return f"{primeiro} {rnd.choice(SOBRENOMES)} {rnd.choice(SOBRENOMES)} FICTICIO"


def _data(rnd, inicio, fim):
    return inicio + timedelta(days=rnd.randint(0, (fim - inicio).days))


def _aa(d):  # data principal: dd/mm/aa
    return d.strftime("%d/%m/%y")


def _aaaa(d):  # coluna (Y2K): dd/mm/aaaa
    return d.strftime("%d/%m/%Y")


def _digitos(rnd, n):
    return "".join(str(rnd.randint(0, 9)) for _ in range(n))


def _cartao(codigo, cert, sufixo):
    # Padrao do arquivo real: 7 + contrato + certificado (6 digitos) + sufixo.
    return f"7{codigo}{cert[-6:]}{sufixo}"


def gerar_empresa(codigo, n_titulares, rnd, com_casos_especiais=False):
    abas = {nome: [] for nome in ("titular", "pis", "mae", "banco", "dependentes")}
    for i in range(1, n_titulares + 1):
        cert = f"{i * 57:07d}"
        subf = f"{rnd.randint(1, 2):03d}"
        matricula = f"{rnd.randint(1, 99999):011d}"
        sexo = rnd.choice(["01", "02"])
        nome = _nome(rnd, sexo)
        nasc = _data(rnd, date(1960, 1, 1), date(2002, 12, 31))
        adm = _data(rnd, date(2005, 1, 1), date(2025, 6, 30))
        vig = max(adm, date(2012, 11, 1))
        plano = rnd.choice(PLANOS)
        # No arquivo real o CPF do titular as vezes chega como numero.
        cpf = int(_digitos(rnd, 11)) if rnd.random() < 0.5 else _digitos(rnd, 11)
        abas["titular"].append([
            "2", subf, cert, nome, matricula, sexo, _aa(nasc), rnd.choice(["1", "2"]),
            cpf, rnd.choice(CARGOS), _aa(adm), _aa(vig), plano, "",
            _aaaa(nasc), _aaaa(adm), _aaaa(vig), "", DATA_VAZIA, "0000", DATA_VAZIA, "",
        ])
        if rnd.random() < 0.5:
            abas["pis"].append(["6", subf, cert, nome, matricula, "", "", "", "000", DATA_VAZIA,
                                _digitos(rnd, 11), ""])
        abas["mae"].append(["5", subf, cert, nome, matricula, _cartao(codigo, cert, "001"),
                            _nome(rnd, "02"), "", f"{plano}{vig:%d%m%Y}"])
        banco, agencia = rnd.choice(BANCOS)
        abas["banco"].append(["9", subf, cert, nome, matricula, banco, agencia, "", "",
                              f"{rnd.randint(1, 99999):013d}", str(rnd.randint(0, 9)), "", "", ""])

        for k in range(1, rnd.choice([0, 1, 1, 2, 3]) + 1):
            sexo_dep = rnd.choice(["01", "02"])
            grau = "1" if k == 1 and rnd.random() < 0.5 else "2"  # 1 = conjuge, 2 = filho(a)
            nasc_dep = _data(rnd, date(1965, 1, 1), date(2025, 12, 31))
            vig_dep = max(vig, nasc_dep)
            abas["dependentes"].append([
                "3", subf, cert, f"{k:02d}", _nome(rnd, sexo_dep), _aa(nasc_dep), sexo_dep,
                "2" if grau == "1" else "1", grau, _aa(vig_dep), "",
                _aaaa(nasc_dep), _aaaa(vig_dep), "", DATA_VAZIA, DATA_VAZIA,
                _digitos(rnd, 11), _cartao(codigo, cert, f"{k * 8 + 1:03d}"), _nome(rnd, "02"),
            ])

    if com_casos_especiais:
        # Data impossivel e texto no lugar de data (vao para o Log).
        abas["titular"][0][6] = "31/02/85"
        abas["titular"][1][6] = "SEM DATA"
        # Dependente cujo certificado nao existe na aba de titulares (vai para Pendencias).
        abas["dependentes"].append([
            "3", "001", "0099999", "01", "DEPENDENTE SEM TITULAR FICTICIO", "01/01/15", "01",
            "1", "2", "01/01/15", "", "01/01/2015", "01/01/2015", "", DATA_VAZIA, DATA_VAZIA,
            _digitos(rnd, 11), _cartao(codigo, "0099999", "009"), "MAE FICTICIA DE TESTE",
        ])
        # Titular sem linha na aba NOME MAE (aviso no Log).
        abas["mae"].pop()
    return abas


def salvar(caminho, codigo, abas):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("POS. CADASTRAL (TITULAR)")
    ws.append(CAB_TITULAR_ARQUIVO)
    ws.append(["1", "POSICAO CADASTRAL", f"00{codigo}", _aa(COMPETENCIA), _aa(COMPETENCIA),
               _aaaa(COMPETENCIA), _aaaa(COMPETENCIA), ""])
    ws.append(CAB_TITULAR)
    for linha in abas["titular"]:
        ws.append(linha)

    for titulo, cabecalho, chave in [
        ("POS. CADASTRAL (PIS)", CAB_PIS, "pis"),
        ("POS. CADASTRAL (NOME MÃE)", CAB_NOME_MAE, "mae"),
        ("POS. CADASTRAL (DADOS BANCO)", CAB_BANCO, "banco"),
    ]:
        ws = wb.create_sheet(titulo)
        ws.append(cabecalho)
        for linha in abas[chave]:
            ws.append(linha)

    ws = wb.create_sheet("POS. CADASTRAL (DEPENDENTES)")
    ws.append(CAB_DEPENDENTES)
    for linha in abas["dependentes"]:
        ws.append(linha)
    n_tit, n_banco, n_dep, n_pis = (len(abas[k]) for k in ("titular", "banco", "dependentes", "pis"))
    n_correntistas = n_banco - 1
    total = 1 + n_tit + n_pis + n_banco + n_correntistas + n_dep
    ws.append(CAB_TOTAIS)
    ws.append(["4", str(n_tit), str(n_banco), str(n_correntistas), str(n_dep), str(total), ""])

    wb.save(caminho)


def main():
    rnd = random.Random(20261023)
    entrada = PASTA / "entrada"
    entrada.mkdir(parents=True, exist_ok=True)
    for antigo in entrada.glob("PC9000*.xlsx"):
        antigo.unlink()
    for codigo, empresa, n in EMPRESAS:
        abas = gerar_empresa(codigo, n, rnd, com_casos_especiais=(codigo == "90003"))
        salvar(entrada / f"PC{codigo} 10.2026.xlsx", codigo, abas)
    linhas = ["EMPRESA;COD_SAUDE;COD_DENTAL"] + [f"{e};{c};" for c, e, _ in EMPRESAS]
    (PASTA / "COD_EMPRESAS_EXEMPLO.csv").write_text("\n".join(linhas) + "\n", encoding="utf-8")
    print(f"3 arquivos gerados em {entrada}")


if __name__ == "__main__":
    main()
