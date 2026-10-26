"""Achados da revisao por 5 agentes (rodada 2, 2026-10-23); ids C1..C15 do consenso."""
import pytest

from consolidacao.erros import ErroValidacao
from tests.fabrica import DEPENDENTE_OK, SUPLEMENTAR_OK, TITULAR_OK, criar_arquivo_pc


def _arquivo_ok(pasta, nome="PC123 x.xlsx"):
    p = pasta / nome
    criar_arquivo_pc(p, TITULAR_OK, DEPENDENTE_OK, SUPLEMENTAR_OK)
    return p


# C1 ------------------------------------------------------------------------

def test_c1_saida_igual_a_entrada_e_recusada_sem_tocar_no_arquivo(tmp_path):
    from consolidacao.pipeline import executar

    a = _arquivo_ok(tmp_path)
    original = a.read_bytes()
    # Mesmo arquivo escrito de outro jeito (caixa diferente, caminho com "..").
    outro_nome = tmp_path / "sub" / ".." / a.name.upper()
    (tmp_path / "sub").mkdir()
    for saida in (a, outro_nome):
        with pytest.raises(ErroValidacao, match="e um dos arquivos de entrada"):
            executar([a], {"123": "A"}, saida)
    assert a.read_bytes() == original


def test_c1_saida_igual_ao_csv_de_empresas_e_recusada(tmp_path):
    from consolidacao.pipeline import executar

    a = _arquivo_ok(tmp_path)
    csv = tmp_path / "e.csv"
    csv.write_bytes(b"EMPRESA;COD_SAUDE\nA;123\n")
    with pytest.raises(ErroValidacao, match="e o arquivo de empresas"):
        executar([a], {"123": "A"}, csv, arquivo_empresas=csv)
    assert csv.read_bytes() == b"EMPRESA;COD_SAUDE\nA;123\n"


def test_c1_cli_recusa_saida_igual_a_entrada(tmp_path, capsys):
    from tests.fabrica import escrever_csv_empresas
    from tratamento import main

    a = _arquivo_ok(tmp_path)
    escrever_csv_empresas(tmp_path / "e.csv", [("ACME", "123")])
    codigo = main(["--entrada", str(tmp_path), "--empresas", str(tmp_path / "e.csv"),
                   "--saida", str(a)])
    assert codigo == 1
    assert "e um dos arquivos de entrada" in capsys.readouterr().out


# C4 ------------------------------------------------------------------------

def test_c4_csv_em_codificacao_desconhecida_vira_erro_de_validacao(tmp_path):
    from consolidacao.empresas import carregar_mapa_empresas

    p = tmp_path / "e.csv"
    # 0x81 nao existe no cp1252 nem e UTF-8 valido.
    p.write_bytes(b"EMPRESA;COD_SAUDE\nA\x81;1\n")
    with pytest.raises(ErroValidacao, match="codificacao"):
        carregar_mapa_empresas(p)


def test_c4_csv_que_e_pasta_vira_erro_de_validacao(tmp_path):
    from consolidacao.empresas import carregar_mapa_empresas

    with pytest.raises(ErroValidacao, match="Nao foi possivel ler o arquivo de empresas"):
        carregar_mapa_empresas(tmp_path)


def test_c4_csv_sem_permissao_vira_erro_de_validacao(tmp_path, monkeypatch):
    from pathlib import Path

    from consolidacao.empresas import carregar_mapa_empresas

    p = tmp_path / "e.csv"
    p.write_bytes(b"EMPRESA;COD_SAUDE\nA;1\n")

    def travado(self):
        raise PermissionError(13, "Permission denied", str(self))

    monkeypatch.setattr(Path, "read_bytes", travado)
    with pytest.raises(ErroValidacao, match="em uso ou sem permissao"):
        carregar_mapa_empresas(p)


# C5 ------------------------------------------------------------------------

def _consolidado(base, pendencias=None, log=None):
    import pandas as pd

    from consolidacao.processamento import Consolidado

    vazio = pd.DataFrame({"X": pd.Series([], dtype=object)})
    return Consolidado(
        consolidado=base,
        pendencias=vazio if pendencias is None else pendencias,
        log=vazio if log is None else log,
    )


@pytest.mark.parametrize("aba", ["base", "pendencias"])
def test_c5_aba_acima_do_limite_do_excel_e_recusada_antes_de_gravar(tmp_path, aba):
    import pandas as pd

    from consolidacao.exportacao import escrever_excel

    grande = pd.DataFrame({"X": pd.Series([None] * 1_048_576, dtype=object)})
    pequeno = pd.DataFrame({"X": pd.Series(["a"], dtype=object)})
    c = _consolidado(grande, pequeno) if aba == "base" else _consolidado(pequeno, grande)
    destino = tmp_path / "saida.xlsx"
    with pytest.raises(ErroValidacao, match="1.048.575"):
        escrever_excel(c, destino)
    assert list(tmp_path.iterdir()) == []


def test_c5_aba_no_limite_exato_e_aceita():
    import pandas as pd

    from consolidacao.exportacao import validar_limite_linhas

    no_limite = pd.DataFrame({"X": pd.Series([None] * 1_048_575, dtype=object)})
    validar_limite_linhas(_consolidado(no_limite))


# C6 ------------------------------------------------------------------------

def test_c6_temporario_comeca_com_til_cifrao_e_some_em_qualquer_erro(tmp_path, monkeypatch):
    from datetime import date
    from pathlib import Path

    import openpyxl

    from consolidacao.exportacao import escrever_excel
    from consolidacao.processamento import listar_arquivos_entrada

    vistos = []

    def falhar(self, nome):
        vistos.append(Path(nome))
        Path(nome).write_bytes(b"meio arquivo")
        raise RuntimeError("falha inesperada do openpyxl")

    monkeypatch.setattr(openpyxl.Workbook, "save", falhar)
    pasta = tmp_path / "saida"
    with pytest.raises(RuntimeError):
        escrever_excel(_consolidado(_tabela()), pasta / "r.xlsx")
    monkeypatch.undo()
    assert vistos and vistos[0].name.startswith("~$")
    assert list(pasta.iterdir()) == []
    # Um temporario esquecido (ex.: processo morto) nao e lido como entrada no proximo run.
    _arquivo_ok(pasta)
    (pasta / vistos[0].name).write_bytes(b"x")
    assert [p.name for p in listar_arquivos_entrada(pasta, date(2026, 10, 23))] == ["PC123 x.xlsx"]


def _tabela():
    import pandas as pd

    return pd.DataFrame({"X": pd.Series(["a"], dtype=object)})


def test_c6_progresso_avisa_a_gravacao_no_fim(tmp_path):
    from consolidacao.pipeline import ETAPA_GRAVACAO, executar

    a, b = _arquivo_ok(tmp_path, "PC123 a.xlsx"), _arquivo_ok(tmp_path, "PC123 b.xlsx")
    vistos = []
    executar([a, b], {"123": "A"}, tmp_path / "o" / "r.xlsx",
             progresso=lambda i, total, nome: vistos.append((i, total, nome)))
    assert vistos == [(1, 2, "PC123 a.xlsx"), (2, 2, "PC123 b.xlsx"), (2, 2, ETAPA_GRAVACAO)]
    assert ETAPA_GRAVACAO == "Gravando arquivo de saida"


def test_c6_erro_no_callback_de_progresso_nao_aborta(tmp_path):
    from consolidacao.pipeline import executar

    def quebrado(i, total, nome):
        raise ValueError("bug na barra de progresso")

    saida = tmp_path / "o" / "r.xlsx"
    c = executar([_arquivo_ok(tmp_path)], {"123": "A"}, saida, progresso=quebrado)
    assert saida.exists() and len(c.consolidado) == 4


def test_c6_cancelar_pelo_callback_de_progresso_continua_valendo(tmp_path):
    from consolidacao.erros import ProcessamentoCancelado
    from consolidacao.pipeline import executar

    def cancela(i, total, nome):
        raise ProcessamentoCancelado("cancelado")

    saida = tmp_path / "o" / "r.xlsx"
    with pytest.raises(ProcessamentoCancelado):
        executar([_arquivo_ok(tmp_path)], {"123": "A"}, saida, progresso=cancela)
    assert not saida.exists()


def test_c6_excecao_inesperada_vira_erro_de_validacao(tmp_path, monkeypatch):
    import consolidacao.pipeline as pipeline

    def explode(c, destino):
        raise KeyError("coluna")

    monkeypatch.setattr(pipeline, "escrever_excel", explode)
    with pytest.raises(ErroValidacao, match="Erro inesperado") as info:
        pipeline.executar([_arquivo_ok(tmp_path)], {"123": "A"}, tmp_path / "r.xlsx")
    assert isinstance(info.value.__cause__, KeyError)


@pytest.mark.parametrize("excecao", [MemoryError, KeyboardInterrupt])
def test_c6_memoria_e_interrupcao_nao_sao_engolidas(tmp_path, monkeypatch, excecao):
    import consolidacao.pipeline as pipeline

    def explode(c, destino):
        raise excecao()

    monkeypatch.setattr(pipeline, "escrever_excel", explode)
    with pytest.raises(excecao):
        pipeline.executar([_arquivo_ok(tmp_path)], {"123": "A"}, tmp_path / "r.xlsx")


# C7 ------------------------------------------------------------------------

@pytest.mark.parametrize("conteudo", [b"isto nao e um xlsx", b"", b"PK\x03\x04truncado"])
def test_c7_arquivo_corrompido_tem_mensagem_em_portugues(tmp_path, conteudo):
    from consolidacao.processamento import processar_arquivos

    ruim = tmp_path / "PC123 ruim.xlsx"
    ruim.write_bytes(conteudo)
    c = processar_arquivos([ruim, _arquivo_ok(tmp_path)], {"123": "A"})
    assert c.log["STATUS"].tolist() == ["ERRO", "OK"]
    assert c.log["MENSAGEM"][0] == "Nao foi possivel abrir o arquivo: arquivo corrompido ou nao e .xlsx."
    assert len(c.consolidado) == 4


def test_c7_arquivo_em_uso_tem_mensagem_em_portugues(tmp_path, monkeypatch):
    import consolidacao.processamento as processamento

    def em_uso(caminho):
        raise PermissionError(13, "Permission denied", str(caminho))

    monkeypatch.setattr(processamento, "abrir_planilha", em_uso)
    r = processamento.processar_arquivo(_arquivo_ok(tmp_path), {"123": "A"})
    assert (r.status, r.mensagem) == (
        "ERRO", "Nao foi possivel abrir o arquivo: arquivo em uso ou sem permissao."
    )


# C2 ------------------------------------------------------------------------

from tests import fabrica  # noqa: E402

VARIANTES_SEM_VALOR = [
    fabrica.FORMULA_TEXTO_CACHE_VAZIO,
    fabrica.FORMULA_TEXTO_SEM_CACHE,
    fabrica.FORMULA_NUMERO_SEM_CACHE,
    fabrica.NUMERO_V_VAZIO,
    fabrica.LOGICO_V_VAZIO,
]
IDS_SEM_VALOR = ["formula_cache_vazio", "formula_sem_cache", "formula_numero", "v_vazio", "logico_v_vazio"]


def _planilha_com_marcador(caminho, linhas, variante):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "A"
    for linha in linhas:
        ws.append(linha)
    wb.save(caminho)
    fabrica.trocar_celula(caminho, "__E__", variante)


@pytest.mark.parametrize("variante", VARIANTES_SEM_VALOR, ids=IDS_SEM_VALOR)
def test_c2_celula_com_formula_ou_v_vazio_conta_na_faixa(tmp_path, variante):
    from consolidacao.leitura import ler_aba

    p = tmp_path / "x.xlsx"
    _planilha_com_marcador(p, [["NUMERO DO CERTIFICADO", "NOME"], ["001", "ANA"], ["__E__"]], variante)
    df = ler_aba(p, "A")
    # Como o readxl (rev2/A_opus_max/gen_ler3.py): 2 linhas, a segunda toda NA.
    assert df["NUMERO DO CERTIFICADO"].tolist() == ["001", None]
    assert df["NOME"].tolist() == ["ANA", None]


@pytest.mark.parametrize("variante", VARIANTES_SEM_VALOR, ids=IDS_SEM_VALOR)
def test_c2_celula_sem_valor_conta_na_faixa_de_colunas(tmp_path, variante):
    from consolidacao.leitura import ler_aba

    p = tmp_path / "x.xlsx"
    _planilha_com_marcador(p, [["A", "B"], ["1", "2", "__E__"]], variante)
    # O R rejeita a aba (3 colunas x 2 col_types); aqui a coluna sem nome e
    # rejeitada adiante, como a coluna orfa com valor.
    assert list(ler_aba(p, "A").columns) == ["A", "B", ""]


def test_c2_dependente_com_totais_em_formula_sem_cache_nao_perde_registros(tmp_path):
    import openpyxl

    from consolidacao.leitura import ler_dependente

    # rev2/B/dep.py (d01): 3 dependentes + 2 linhas de resumo gravadas por biblioteca.
    cab = ["NUMERO DO CERTIFICADO", "NOME DEPENDENTE", "CODIGO DO DEPENDENTE", "GRAU PARENTESCO",
           "NOME DA MAE DO DEPENDENTE", "NUMERO DO CARTAO DO DEPENDENTE"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(cab)
    for i in range(1, 4):
        ws.append([f"C{i}", f"DEP{i}", "01", "FILHO", "MAE", f"K{i}"])
    ws.append(["=COUNTA(A2:A4)"])
    ws.append(['="TOTAL: "&COUNTA(B2:B4)'])
    p = tmp_path / "d01.xlsx"
    wb.save(p)
    assert ler_dependente(p, "S", "REF").dados["NOME DEPENDENTE"].tolist() == ["DEP1", "DEP2", "DEP3"]


def test_c2_titular_com_formula_vazia_abaixo_dos_dados_e_rejeitado_como_no_r(tmp_path):
    import openpyxl

    from consolidacao.processamento import processar_arquivo

    # rev2/A_opus_max/gen_e2e1.py (E7): o R para com "Ha titular sem NUMERO DO CERTIFICADO".
    p = _arquivo_ok(tmp_path)
    wb = openpyxl.load_workbook(p)
    ws = wb["TITULAR"]
    ws.cell(ws.max_row + 1, 1, "__E__")
    wb.save(p)
    fabrica.trocar_celula(p, "__E__", fabrica.FORMULA_TEXTO_CACHE_VAZIO)
    r = processar_arquivo(p, {"123": "A"})
    assert r.status == "ERRO"
    assert "Ha titular sem NUMERO DO CERTIFICADO" in r.mensagem


def test_c2_varredura_le_celulas_sem_atributo_r(tmp_path):
    import zipfile

    from consolidacao.leitura import varrer_celulas

    p = tmp_path / "x.xlsx"
    _planilha_com_marcador(p, [["A"]], fabrica.FORMULA_TEXTO_SEM_CACHE)
    xml = (
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
        '<row><c t="inlineStr"><is><t>A</t></is></c><c><v>7</v></c></row>'
        '<row r="4"><c r="C4"><f>1+1</f></c><c><v/></c></row>'
        "</sheetData></worksheet>"
    )
    q = tmp_path / "y.xlsx"
    with zipfile.ZipFile(p) as origem, zipfile.ZipFile(q, "w") as destino:
        for item in origem.infolist():
            dados = origem.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                dados = xml.encode()
            destino.writestr(item, dados)
    assert list(varrer_celulas(q, "A")) == [
        (1, 1, "inlineStr", None, False),
        (1, 2, "n", "7", False),
        (4, 3, "n", None, True),
        (4, 4, "n", "", False),
    ]


# C11 -----------------------------------------------------------------------

def test_c11_titular_com_linha_em_branco_entre_registro_e_cabecalho_e_rejeitado(tmp_path):
    import openpyxl

    from consolidacao.processamento import processar_arquivo

    # rev2/A_opus_max/gen_e2e1.py (E6): o R para na leitura do cabecalho.
    p = _arquivo_ok(tmp_path)
    wb = openpyxl.load_workbook(p)
    ws = wb["TITULAR"]
    ws.cell(2, 1, "LINHA 2")
    ws.insert_rows(3)
    wb.save(p)
    r = processar_arquivo(p, {"123": "A"})
    assert (r.status, r.mensagem) == ("ERRO", "A aba 'TITULAR' nao tem cabecalhos identificaveis.")


@pytest.mark.parametrize("linhas, esperado", [
    # Conteudo antes do skip: o cabecalho e a linha do skip, mesmo vazia.
    ([["X"], [], [], ["A"], ["1"]], "erro"),
    # Nada antes do skip: linhas vazias depois dele sao puladas.
    ([[], [], [], ["A"], ["1"]], ["1"]),
    ([["X"], ["Y"], ["A"], ["1"]], ["1"]),
])
def test_c11_inicio_e_o_maior_entre_skip_e_primeira_linha_com_conteudo(tmp_path, linhas, esperado):
    import openpyxl

    from consolidacao.leitura import ler_aba

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "T"
    for i, linha in enumerate(linhas, start=1):
        for j, v in enumerate(linha, start=1):
            ws.cell(i, j, v)
    p = tmp_path / "x.xlsx"
    wb.save(p)
    if esperado == "erro":
        with pytest.raises(ErroValidacao, match="A aba 'T' nao tem cabecalhos identificaveis."):
            ler_aba(p, "T", skip=2)
    else:
        assert ler_aba(p, "T", skip=2)["A"].tolist() == esperado


# C8 ------------------------------------------------------------------------

def test_c8_celula_so_com_quebra_de_linha_vira_na_e_o_resto_fica(tmp_path):
    import openpyxl

    from consolidacao.leitura import ler_aba

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "A"
    ws.append(["NOME DA MAE"])
    for v in ["\n", " \t\n ", "\r\n", "\nabc\n", "\xa0", " x "]:
        ws.append([v])
    p = tmp_path / "x.xlsx"
    wb.save(p)
    # readxl (rev2/B c07_espacos): so-quebra vira NA; "\nabc\n" e NBSP ficam.
    assert ler_aba(p, "A")["NOME DA MAE"].tolist() == [None, None, None, "\nabc\n", "\xa0", "x"]


# C9 ------------------------------------------------------------------------

@pytest.mark.parametrize("valor", [
    "２０２０-01-15",  # digitos de largura total
    "٢٠٢٠-01-15",  # digitos arabes
    "\xa005/05/1980",  # NBSP antes do campo
    "05/ 05/1980",  # espaco Unicode entre os campos
])
def test_c9_digitos_e_espacos_unicode_nao_valem_em_data(valor):
    from consolidacao.datas import converter_celula_data

    data, obs = converter_celula_data(valor, "DATA", 1, "C")
    # strptime do R: data invalida, vai para o log (rev2/B strs.R, rev2/A datas_r.R).
    assert data is None and obs is not None


def test_c9_espaco_ascii_continua_aceito():
    from datetime import date

    from consolidacao.datas import parse_formato_r

    assert parse_formato_r("05/ 05/1980", "%d/%m/%Y") == date(1980, 5, 5)


# C10 -----------------------------------------------------------------------

def test_c10_pasta_1904_data_em_coluna_texto_vira_serial_1904(tmp_path):
    from datetime import date, datetime

    import openpyxl
    from openpyxl.utils.datetime import CALENDAR_MAC_1904

    from consolidacao.leitura import ler_aba

    wb = openpyxl.Workbook()
    wb.epoch = CALENDAR_MAC_1904
    ws = wb.active
    ws.title = "A"
    ws.append(["CODIGO", "DATA DE NASCIMENTO"])
    ws.append([datetime(2020, 1, 15), datetime(2020, 1, 15)])
    ws.append([datetime(1950, 6, 1), datetime(1950, 6, 1)])
    p = tmp_path / "x.xlsx"
    wb.save(p)
    df = ler_aba(p, "A")
    # readxl devolve o serial gravado (rev2/B c05_1904): R "42383" e "16953".
    assert df["CODIGO"].tolist() == ["42383", "16953"]
    assert [v.date() for v in df["DATA DE NASCIMENTO"]] == [date(2020, 1, 15), date(1950, 6, 1)]


# C3 ------------------------------------------------------------------------

DEPENDENTES_COM_TIPO = [dict({"TIPO DE REGISTRO": "3"}, **d) for d in DEPENDENTE_OK]


def _arquivo_dependente_com_tipo(pasta, fim):
    """Aba DEPENDENTE no formato real (coluna TIPO DE REGISTRO), com as linhas
    finais trocadas por `fim`."""
    import openpyxl

    p = pasta / "PC123 x.xlsx"
    criar_arquivo_pc(p, TITULAR_OK, DEPENDENTES_COM_TIPO, SUPLEMENTAR_OK)
    wb = openpyxl.load_workbook(p)
    ws = wb["DEPENDENTE"]
    ws.delete_rows(ws.max_row - 1, 2)
    for linha in fim:
        ws.append(linha)
    wb.save(p)
    return p


TOTAIS_REAIS = [["TIPO DE REGISTRO", "TOTAL DE TITULARES", "TOTAL DE DEPENDENTES"], ["4", "2", "2"]]


@pytest.mark.parametrize("tipo", ["4", " 4 ", 4])
def test_c3_formato_real_remove_os_totais_sem_aviso(tmp_path, tipo):
    from consolidacao.leitura import ler_dependente

    fim = [TOTAIS_REAIS[0], [tipo, "2", "2"]]
    r = ler_dependente(_arquivo_dependente_com_tipo(tmp_path, fim), "DEPENDENTE", "PC1")
    assert r.dados["NOME DEPENDENTE"].tolist() == ["CAIO TESTE", "DORA SEM TITULAR"]
    assert r.avisos_estrutura == []


def test_c3_sem_linha_tipo_4_nao_corta_e_avisa(tmp_path):
    from consolidacao.processamento import processar_arquivo

    # Exportacao truncada: a aba termina no ultimo dependente, sem os totais.
    p = _arquivo_dependente_com_tipo(tmp_path, [])
    r = processar_arquivo(p, {"123": "A"})
    assert r.status == "OK" and r.n_dependentes == 2
    assert "DEPENDENTE: ultima linha sem TIPO DE REGISTRO 4 (resumo); nenhuma linha removida" in r.avisos


def test_c3_um_dependente_sem_totais_nao_vira_erro(tmp_path):
    import openpyxl

    from consolidacao.leitura import ler_dependente

    p = _arquivo_dependente_com_tipo(tmp_path, [])
    wb = openpyxl.load_workbook(p)
    wb["DEPENDENTE"].delete_rows(3)
    wb.save(p)
    r = ler_dependente(p, "DEPENDENTE", "PC1")
    assert r.dados["NOME DEPENDENTE"].tolist() == ["CAIO TESTE"]
    assert len(r.avisos_estrutura) == 1


def test_c3_sem_coluna_tipo_corta_como_o_r(tmp_path):
    from consolidacao.leitura import ler_dependente

    r = ler_dependente(_arquivo_ok(tmp_path), "DEPENDENTE", "PC1")
    assert len(r.dados) == 2 and r.avisos_estrutura == []


# C12 -----------------------------------------------------------------------

def test_c12_linha_vazia_no_meio_dos_dependentes_fica_e_e_contada_no_log(tmp_path):
    from consolidacao.processamento import processar_arquivo

    p = tmp_path / "PC123 x.xlsx"
    criar_arquivo_pc(p, TITULAR_OK, [DEPENDENTE_OK[0], {}, {}, DEPENDENTE_OK[1]], SUPLEMENTAR_OK)
    r = processar_arquivo(p, {"123": "A"})
    # Paridade: o R mantem as linhas vazias como registros (e pendencias).
    assert r.status == "OK" and r.n_dependentes == 4 and len(r.pendencias) == 3
    assert r.avisos.endswith("2 linha(s) vazia(s) em DEPENDENTES")


def test_c12_sem_linha_vazia_sem_aviso(tmp_path):
    from consolidacao.processamento import processar_arquivo

    r = processar_arquivo(_arquivo_ok(tmp_path), {"123": "A"})
    assert "vazia" not in r.avisos


# C13 -----------------------------------------------------------------------

def test_c13_mesmo_codigo_de_empresa_em_dois_arquivos_avisa_nos_dois(tmp_path):
    from consolidacao.processamento import processar_arquivos

    a, b = _arquivo_ok(tmp_path, "PC123 a.xlsx"), _arquivo_ok(tmp_path, "PC123 b v2.xlsx")
    outro = _arquivo_ok(tmp_path, "PC456 c.xlsx")
    c = processar_arquivos([a, b, outro], {"123": "A", "456": "B"})
    avisos = c.log["AVISOS_ESTRUTURA"].tolist()
    assert avisos[0].endswith("CODIGO_EMPRESA '123' repetido em outro(s) arquivo(s): PC123 b v2.xlsx")
    assert avisos[1].endswith("CODIGO_EMPRESA '123' repetido em outro(s) arquivo(s): PC123 a.xlsx")
    assert "repetid" not in avisos[2]
    # Aviso, nao erro: os dois entram na base.
    assert c.log["STATUS"].tolist() == ["OK", "OK", "OK"] and len(c.consolidado) == 12


def test_c13_mesma_referencia_em_pastas_diferentes(tmp_path):
    from consolidacao.processamento import processar_arquivos

    (tmp_path / "x").mkdir()
    (tmp_path / "y").mkdir()
    a, b = _arquivo_ok(tmp_path / "x", "PC123 m.xlsx"), _arquivo_ok(tmp_path / "y", "PC123 m.xlsx")
    c = processar_arquivos([a, b], {"123": "A"})
    assert c.log["AVISOS_ESTRUTURA"][0].endswith(
        "REFERENCIA 'PC123 m' repetida em outro(s) arquivo(s): PC123 m.xlsx"
    )


def test_c13_arquivo_com_erro_nao_conta_como_repetido(tmp_path):
    from consolidacao.processamento import processar_arquivos

    ruim = tmp_path / "PC123 ruim.xlsx"
    ruim.write_bytes(b"x")
    c = processar_arquivos([_arquivo_ok(tmp_path), ruim], {"123": "A"})
    assert "repetid" not in c.log["AVISOS_ESTRUTURA"][0]


def test_c13_paridade_ignora_so_os_avisos_que_o_r_nao_tem():
    from tests.paridade.comparar import sem_avisos_so_python

    assert sem_avisos_so_python(
        "Titulares sem CARTAO apos a juncao: 1 | CODIGO_EMPRESA '1' repetido em outro(s) "
        "arquivo(s): a.xlsx | 2 linha(s) vazia(s) em DEPENDENTES"
    ) == "Titulares sem CARTAO apos a juncao: 1"
    assert sem_avisos_so_python("REFERENCIA 'X' repetida em outro(s) arquivo(s): a.xlsx") is None
    assert sem_avisos_so_python(None) is None
    assert sem_avisos_so_python("Titulares sem CARTAO apos a juncao: 1") == (
        "Titulares sem CARTAO apos a juncao: 1"
    )


# C14 -----------------------------------------------------------------------

def test_c14_celula_com_erro_do_excel_vira_none(tmp_path):
    from consolidacao.leitura import ler_aba

    p = tmp_path / "x.xlsx"
    _planilha_com_marcador(p, [["A", "B"], ["1", "__E__"], ["2", "x"]],
                           lambda ref: f'<c r="{ref}" t="e"><v>#N/A</v></c>')
    assert ler_aba(p, "A")["B"].tolist() == [None, "x"]


def test_c14_suplementar_sem_certificado_e_erro(tmp_path):
    from consolidacao.leitura import ler_suplementar
    from tests.fabrica import ABA_SUPLEMENTAR

    sup = SUPLEMENTAR_OK + [{"NUMERO DO CERTIFICADO": None, "CARTÃO": "C9", "NOME DA MÃE": "M"}]
    p = tmp_path / "PC1.xlsx"
    criar_arquivo_pc(p, TITULAR_OK, DEPENDENTE_OK, sup)
    with pytest.raises(ErroValidacao, match=r"\[POS. CADASTRAL \(NOME M E\): PC1\] Ha registro suplementar sem NUMERO DO CERTIFICADO."):
        ler_suplementar(p, ABA_SUPLEMENTAR, "PC1")


@pytest.mark.parametrize("valor", [None, " "])
def test_c14_aba_suplementar_sem_registros_e_erro(tmp_path, valor):
    from consolidacao.leitura import ler_suplementar
    from tests.fabrica import ABA_SUPLEMENTAR

    sup = [{"NUMERO DO CERTIFICADO": valor, "CARTÃO": valor, "NOME DA MÃE": valor}]
    p = tmp_path / "PC1.xlsx"
    criar_arquivo_pc(p, TITULAR_OK, DEPENDENTE_OK, sup)
    with pytest.raises(ErroValidacao, match="A aba suplementar nao contem registros."):
        ler_suplementar(p, ABA_SUPLEMENTAR, "PC1")


def test_c14_duas_colunas_de_data_para_a_mesma_y2k_e_erro():
    import pandas as pd

    from consolidacao.datas import validar_colunas_data

    df = pd.DataFrame({
        "DATA DE NASCIMENTO": ["01/01/1990"],
        "DATA DE NASCIMENTO.": ["01/01/1990"],
        "DATA DE NASCIMENTO (Y2K)": ["01/01/1990"],
    }, dtype=object)
    with pytest.raises(ErroValidacao, match=r"\[C\] Mais de uma coluna de data corresponde a 'DATA DE NASCIMENTO \(Y2K\)'."):
        validar_colunas_data(df, "C")


def test_c14_csv_sem_pares_validos_e_erro(tmp_path):
    from consolidacao.empresas import carregar_mapa_empresas

    p = tmp_path / "e.csv"
    p.write_bytes(b"EMPRESA;COD_SAUDE\n;1\nNA;2\nA;\n")
    with pytest.raises(ErroValidacao, match="O COD_EMPRESAS.csv nao tem pares EMPRESA/COD_SAUDE validos."):
        carregar_mapa_empresas(p)


def _arquivo_subf(pasta):
    p = pasta / "PC123 subf.xlsx"
    tit = [dict(TITULAR_OK[0], **{"NUMERO DA SUBFATURA": "0002"}),
           dict(TITULAR_OK[1], **{"NUMERO DA SUBFATURA": "0001"})]
    dep = [dict(d, **{"NUMERO DA SUBFATURA": "0009"}) for d in DEPENDENTE_OK]
    criar_arquivo_pc(p, tit, dep, SUPLEMENTAR_OK)
    return p


def test_c14_subfatura_do_dependente_e_fallback_do_titular(tmp_path):
    from consolidacao.processamento import processar_arquivo

    r = processar_arquivo(_arquivo_subf(tmp_path), {"123": "A"})
    dep = r.final[r.final["TITULARIDADE"] == "DEPENDENTE"]
    # Com titular, vale a do titular; sem titular, a do proprio dependente (coalesce do R).
    assert dep["SubF"].tolist() == ["0002", "0009"]
    assert r.pendencias["NUMERO DA SUBFATURA"].tolist() == ["0009"]
    assert ".SUBF_DEPENDENTE" not in r.final.columns


def test_c14_base_ordenada_por_referencia_subf_e_certificado(tmp_path):
    from consolidacao.processamento import processar_arquivos

    c = processar_arquivos([_arquivo_subf(tmp_path)], {"123": "A"})
    assert c.consolidado["SubF"].tolist() == ["0001", "0002", "0002", "0009"]
    assert c.consolidado["NUMERO DO CERTIFICADO"].tolist() == ["002", "001", "001", "999"]


def test_c14_aba_sem_colunas_nao_e_exportada():
    import openpyxl
    import pandas as pd

    from consolidacao.exportacao import adicionar_aba_formatada

    with pytest.raises(ValueError, match="A aba 'X' nao tem colunas para exportar."):
        adicionar_aba_formatada(openpyxl.Workbook(), pd.DataFrame(), "X", "T")


@pytest.mark.parametrize("x, texto", [
    # Valores conferidos com format(x, scientific = FALSE, trim = TRUE) do R 4.6.1.
    (1234.5678, "1234.568"), (0.000123456789, "0.0001234568"), (12345678.9, "12345679"),
    (5.0, "5"), (0.5, "0.5"), (-1.25, "-1.25"), (1 / 3, "0.3333333"),
    (float("inf"), "Inf"), (float("-inf"), "-Inf"), (float("nan"), "NaN"),
    (99999999.98999999, "100000000"),
])
def test_c14_formatar_numero_como_o_r(x, texto):
    from consolidacao.datas import _formatar_numero_r

    assert _formatar_numero_r(x) == texto


def test_c14_logico_vira_true_false():
    from consolidacao.datas import texto_celula
    from consolidacao.leitura import _como_texto

    assert [texto_celula(True), texto_celula(False)] == ["TRUE", "FALSE"]
    assert [_como_texto(True), _como_texto(False)] == ["TRUE", "FALSE"]
