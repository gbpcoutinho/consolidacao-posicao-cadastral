"""Gera arquivos sinteticos (dados falsos) para os testes."""
import re
import shutil
import zipfile
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill

ABA_SUPLEMENTAR = "POS. CADASTRAL (NOME MÃE)"

TITULAR_OK = [
    {
        "TIPO DE REGISTRO": "1",
        "NUMERO DA SUBFATURA": "0001",
        "NUMERO DO CERTIFICADO": "001",
        "NOME DO SEGURADO": "ANA TESTE",
        "PLANO": "P1",
        "DATA DE NASCIMENTO": "05/05/80",
        "REGIÃO": "SUL",
        "SEXO DO SEGURADO": "F",
        "NUMERO DO CPF": "00011122233",
    },
    {
        "TIPO DE REGISTRO": "1",
        "NUMERO DA SUBFATURA": "0001",
        "NUMERO DO CERTIFICADO": "002",
        "NOME DO SEGURADO": "BRUNO TESTE",
        "PLANO": "P2",
        "DATA DE NASCIMENTO": datetime(1975, 3, 1),
        "REGIÃO": "NORTE",
        "SEXO DO SEGURADO": "M",
        "NUMERO DO CPF": "00044455566",
    },
]

DEPENDENTE_OK = [
    {
        "NUMERO DO CERTIFICADO": "001",
        "CODIGO DO DEPENDENTE": "01",
        "NOME DEPENDENTE": "CAIO TESTE",
        "GRAU PARENTESCO": "FILHO",
        "NOME DA MÃE DO DEPENDENTE": "ANA TESTE",
        "CARTÃO DO DEPENDENTE": "C001-01",
        "SEXO DO DEPENDENTE": "M",
        "DATA DE NASCIMENTO": "10/10/2010",
    },
    {
        "NUMERO DO CERTIFICADO": "999",
        "CODIGO DO DEPENDENTE": "02",
        "NOME DEPENDENTE": "DORA SEM TITULAR",
        "GRAU PARENTESCO": "FILHA",
        "NOME DA MÃE DO DEPENDENTE": "MARIA TESTE",
        "CARTÃO DO DEPENDENTE": "C999-02",
        "SEXO DO DEPENDENTE": "F",
        "DATA DE NASCIMENTO": "01/02/2012",
    },
]

SUPLEMENTAR_OK = [
    {"NUMERO DO CERTIFICADO": "001", "CARTÃO": "C001", "NOME DA MÃE": "MAE DA ANA"},
]


def escrever_csv_empresas(caminho: Path, linhas: list, encoding: str = "utf-8-sig") -> None:
    texto = "EMPRESA;COD_SAUDE;COD_DENTAL\n" + "".join(
        f"{empresa};{codigo};\n" for empresa, codigo in linhas
    )
    Path(caminho).write_bytes(texto.encode(encoding))


def forcar_dimensao(caminho: Path, indice_aba: int, ref: str) -> None:
    """Regrava o <dimension> de uma aba, como fazem alguns geradores de xlsx."""
    caminho = Path(caminho)
    temporario = caminho.with_suffix(".tmp")
    nome_xml = f"xl/worksheets/sheet{indice_aba}.xml"
    with zipfile.ZipFile(caminho) as origem, zipfile.ZipFile(temporario, "w") as destino:
        for item in origem.infolist():
            dados = origem.read(item.filename)
            if item.filename == nome_xml:
                dados = re.sub(rb'<dimension ref="[^"]*"', f'<dimension ref="{ref}"'.encode(), dados)
            destino.writestr(item, dados)
    shutil.move(temporario, caminho)


def _colunas(registros: list, padrao: list) -> list:
    return list((registros or padrao)[0].keys())


def _escrever_registros(ws, registros: list, colunas: list) -> None:
    ws.append(colunas)
    for r in registros:
        ws.append([r.get(c) for c in colunas])


def criar_arquivo_pc(
    caminho: Path,
    titulares: list,
    dependentes: list,
    suplementar: list,
    abas: tuple = ("TITULAR", "DEPENDENTE", ABA_SUPLEMENTAR),
    linhas_vazias_formatadas_no_fim: int = 0,
) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    aba_tit, aba_dep, aba_sup = abas

    ws = wb.create_sheet(aba_tit)
    ws.append(["RELATORIO DE TITULARES"])
    ws.append([])
    _escrever_registros(ws, titulares, _colunas(titulares, TITULAR_OK))

    ws = wb.create_sheet(aba_dep)
    _escrever_registros(ws, dependentes, _colunas(dependentes, DEPENDENTE_OK))
    ws.append(["TOTAL"])
    ws.append([len(dependentes)])
    if linhas_vazias_formatadas_no_fim:
        fill = PatternFill("solid", fgColor="FFFF00")
        inicio = ws.max_row + 1
        for linha in range(inicio, inicio + linhas_vazias_formatadas_no_fim):
            for coluna in range(1, 4):
                ws.cell(row=linha, column=coluna).fill = fill

    ws = wb.create_sheet(aba_sup)
    _escrever_registros(ws, suplementar, _colunas(suplementar, SUPLEMENTAR_OK))

    wb.save(caminho)


def trocar_celula(caminho: Path, marcador: str, xml_celula) -> None:
    """Troca a celula de texto `marcador` pelo XML cru `xml_celula(ref)`, para
    simular geradores que gravam formula sem valor em cache ou <v> vazio."""
    caminho = Path(caminho)
    temporario = caminho.with_suffix(".tmp")
    padrao = re.compile(
        r'<c r="([A-Z]+[0-9]+)"( s="\d+")? t="inlineStr"><is><t>'
        + re.escape(marcador) + r"</t></is></c>"
    )

    def _trocar(m):
        # Mantem o estilo da celula original (ex.: formato de data).
        ref, estilo = m.group(1), m.group(2) or ""
        return xml_celula(ref).replace(f'<c r="{ref}"', f'<c r="{ref}"{estilo}', 1)

    with zipfile.ZipFile(caminho) as origem, zipfile.ZipFile(temporario, "w") as destino:
        for item in origem.infolist():
            dados = origem.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                dados = padrao.sub(_trocar, dados.decode("utf-8")).encode("utf-8")
            destino.writestr(item, dados)
    shutil.move(temporario, caminho)


# Celulas que o readxl conta na faixa lida, mas que o openpyxl le como vazias.
FORMULA_TEXTO_CACHE_VAZIO = lambda ref: f'<c r="{ref}" t="str"><f>""</f><v></v></c>'  # noqa: E731
FORMULA_TEXTO_SEM_CACHE = lambda ref: f'<c r="{ref}" t="str"><f>""</f></c>'  # noqa: E731
FORMULA_NUMERO_SEM_CACHE = lambda ref: f'<c r="{ref}"><f>1+1</f></c>'  # noqa: E731
NUMERO_V_VAZIO = lambda ref: f'<c r="{ref}"><v></v></c>'  # noqa: E731
LOGICO_V_VAZIO = lambda ref: f'<c r="{ref}" t="b"><v></v></c>'  # noqa: E731
