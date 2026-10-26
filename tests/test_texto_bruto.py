"""P11: numero em coluna texto fica com o texto bruto de <v>, como no readxl.

Valores esperados conferidos com o R 4.6.1 (readxl 1.5.0.1) em
rev2/p11/gen.py + rev2/A_opus_max/dump_r.R."""
from datetime import date, datetime

import openpyxl
import pytest
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from tests.fabrica import DEPENDENTE_OK, SUPLEMENTAR_OK, TITULAR_OK, criar_arquivo_pc, trocar_celula


def _numero(v):
    return lambda ref: f'<c r="{ref}"><v>{v}</v></c>'


DATA_ISO = lambda ref: f'<c r="{ref}" t="d"><v>2023-03-15T00:00:00</v></c>'  # noqa: E731


def _planilha(caminho, celula, formato=None, epoca=None):
    """Aba A com as colunas NUMERO DO CERTIFICADO (texto) e DATA X (data); a
    linha 2 recebe o XML cru `celula` nas duas colunas."""
    wb = openpyxl.Workbook()
    if epoca:
        wb.epoch = epoca
    ws = wb.active
    ws.title = "A"
    ws.append(["NUMERO DO CERTIFICADO", "DATA X"])
    ws.append(["__E__", "__E__"])
    if formato:
        for c in ws[2]:
            c.number_format = formato
    wb.save(caminho)
    trocar_celula(caminho, "__E__", celula)


@pytest.mark.parametrize("bruto", [
    "001", "1.2E3", "12345678901234567", "0.3333333333333333", "99999999.98999999",
    "-0", "7.0", "7", " 12 ", "45000.5",
])
def test_numero_em_coluna_texto_fica_com_o_texto_de_v(tmp_path, bruto):
    from consolidacao.leitura import ler_aba

    p = tmp_path / "x.xlsx"
    _planilha(p, _numero(bruto))
    assert ler_aba(p, "A")["NUMERO DO CERTIFICADO"].tolist() == [bruto]


def test_data_iso_em_coluna_texto_fica_como_texto(tmp_path):
    from consolidacao.leitura import ler_aba

    p = tmp_path / "x.xlsx"
    _planilha(p, DATA_ISO)
    df = ler_aba(p, "A")
    assert df["NUMERO DO CERTIFICADO"].tolist() == ["2023-03-15T00:00:00"]
    # Coluna DATA continua convertida.
    assert df["DATA X"].tolist() == [datetime(2023, 3, 15)]


@pytest.mark.parametrize("epoca", [None, CALENDAR_MAC_1904])
@pytest.mark.parametrize("serial", ["60", "45000"])
def test_serial_com_formato_de_data_em_coluna_texto_fica_como_gravado(tmp_path, epoca, serial):
    from consolidacao.leitura import ler_aba

    p = tmp_path / "x.xlsx"
    _planilha(p, _numero(serial), formato="dd/mm/yyyy", epoca=epoca)
    assert ler_aba(p, "A")["NUMERO DO CERTIFICADO"].tolist() == [serial]


def test_coluna_data_continua_numerica(tmp_path):
    from consolidacao.datas import validar_colunas_data
    from consolidacao.leitura import ler_aba

    p = tmp_path / "x.xlsx"
    _planilha(p, _numero("45000"), formato="dd/mm/yyyy")
    df, obs, _ = validar_colunas_data(ler_aba(p, "A"), "C")
    assert df["DATA X"].tolist() == [date(2023, 3, 15)] and obs == []


def test_cabecalho_numerico_usa_o_texto_de_v(tmp_path):
    from consolidacao.leitura import ler_aba

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "A"
    ws.append(["__H__", "B"])
    ws.append(["x", "y"])
    p = tmp_path / "x.xlsx"
    wb.save(p)
    trocar_celula(p, "__H__", _numero("1.50"))
    assert list(ler_aba(p, "A").columns) == ["1.50", "B"]


def test_logico_e_erro_nao_mudam(tmp_path):
    from consolidacao.leitura import ler_aba

    p = tmp_path / "x.xlsx"
    _planilha(p, lambda ref: f'<c r="{ref}" t="b"><v>1</v></c>')
    assert ler_aba(p, "A")["NUMERO DO CERTIFICADO"].tolist() == ["TRUE"]
    _planilha(p, lambda ref: f'<c r="{ref}" t="e"><v>#N/A</v></c>')
    assert ler_aba(p, "A")["NUMERO DO CERTIFICADO"].tolist() == [None]


def test_certificado_numerico_com_zero_a_esquerda_casa_com_as_outras_abas(tmp_path):
    from consolidacao.processamento import processar_arquivo

    # rev2/A_opus_max/gen_e2e3.py: exportador nao-Excel grava <c><v>001</v></c>.
    tit = [dict(TITULAR_OK[0], **{"NUMERO DO CERTIFICADO": "__CERT__"}), TITULAR_OK[1]]
    p = tmp_path / "PC123 x.xlsx"
    criar_arquivo_pc(p, tit, DEPENDENTE_OK, SUPLEMENTAR_OK)
    trocar_celula(p, "__CERT__", _numero("001"))
    r = processar_arquivo(p, {"123": "A"})
    assert r.status == "OK"
    tit001 = r.final[(r.final["TITULARIDADE"] == "TITULAR") & (r.final["NUMERO DO CERTIFICADO"] == "001")]
    assert tit001["NOME DA MAE"].tolist() == ["MAE DA ANA"]
    assert r.pendencias["NUMERO DO CERTIFICADO"].tolist() == ["999"]
