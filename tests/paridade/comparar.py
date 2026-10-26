"""Compara duas planilhas de saida celula a celula (valores e formato de data)."""
import re
from datetime import date, datetime
from pathlib import Path

import openpyxl


def _normalizar(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if v == "":
        return None
    return v


# Avisos que so o Python grava em AVISOS_ESTRUTURA (desvios protetores
# documentados em docs/fase1-conversao.md). Todo o resto do Log e comparado.
AVISOS_SO_PYTHON = [
    re.compile(r"^CODIGO_EMPRESA '.*' repetido em outro\(s\) arquivo\(s\): "),
    re.compile(r"^REFERENCIA '.*' repetida em outro\(s\) arquivo\(s\): "),
    re.compile(r"^[0-9]+ linha\(s\) vazia\(s\) em DEPENDENTES$"),
    re.compile(r"^DEPENDENTE: ultima linha sem TIPO DE REGISTRO 4 \(resumo\); nenhuma linha removida$"),
]


def sem_avisos_so_python(texto):
    if texto is None:
        return None
    itens = [a for a in texto.split(" | ") if not any(p.match(a) for p in AVISOS_SO_PYTHON)]
    return " | ".join(itens) or None


def _linhas(ws):
    # A dimensao gravada pelo openxlsx nao e confiavel; le a aba inteira.
    ws.reset_dimensions()
    return [[c for c in linha] for linha in ws.iter_rows()]


def comparar_planilhas(a: Path, b: Path, max_difs: int = 50) -> list:
    # read_only: arquivos do openxlsx referenciam desenhos que o openpyxl nao acha.
    wa = openpyxl.load_workbook(a, data_only=True, read_only=True)
    wb = openpyxl.load_workbook(b, data_only=True, read_only=True)
    try:
        return _comparar(wa, wb, max_difs)
    finally:
        wa.close()
        wb.close()


def _comparar(wa, wb, max_difs: int) -> list:
    difs = []

    if wa.sheetnames != wb.sheetnames:
        return [f"abas diferentes: {wa.sheetnames} x {wb.sheetnames}"]

    for nome in wa.sheetnames:
        la, lb = _linhas(wa[nome]), _linhas(wb[nome])
        # Linhas totalmente vazias no fim nao contam.
        while la and all(c.value is None for c in la[-1]):
            la.pop()
        while lb and all(c.value is None for c in lb[-1]):
            lb.pop()

        cab_a = [c.value for c in la[0]] if la else []
        cab_b = [c.value for c in lb[0]] if lb else []
        if cab_a != cab_b:
            difs.append(f"[{nome}] cabecalhos diferentes:\n  R : {cab_a}\n  Py: {cab_b}")
            continue
        if len(la) != len(lb):
            difs.append(f"[{nome}] numero de linhas: R={len(la) - 1} Py={len(lb) - 1}")

        for i, (ra, rb) in enumerate(zip(la[1:], lb[1:]), start=2):
            for j, coluna in enumerate(cab_a):
                ca = ra[j] if j < len(ra) else None
                cb = rb[j] if j < len(rb) else None
                va = _normalizar(ca.value if ca is not None else None)
                vb = _normalizar(cb.value if cb is not None else None)
                if nome == "Log" and coluna == "AVISOS_ESTRUTURA":
                    vb = sem_avisos_so_python(vb)
                if va != vb:
                    difs.append(f"[{nome}] linha {i} coluna '{coluna}': R={va!r} Py={vb!r}")
                elif isinstance(va, date) and ca.number_format != cb.number_format:
                    difs.append(
                        f"[{nome}] linha {i} coluna '{coluna}': formato "
                        f"R={ca.number_format!r} Py={cb.number_format!r}"
                    )
                if len(difs) >= max_difs:
                    return difs
    return difs
