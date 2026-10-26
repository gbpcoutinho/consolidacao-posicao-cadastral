"""Gera um conjunto de planilhas sinteticas que cobre as regras do tratamento."""
from datetime import datetime
from pathlib import Path

import openpyxl

from tests.fabrica import (
    FORMULA_TEXTO_CACHE_VAZIO,
    DEPENDENTE_OK,
    SUPLEMENTAR_OK,
    TITULAR_OK,
    criar_arquivo_pc,
    escrever_csv_empresas,
    forcar_dimensao,
    trocar_celula,
)


def _com(registro: dict, **valores) -> dict:
    return {**registro, **valores}


def gerar(pasta: Path) -> Path:
    pasta = Path(pasta)
    pasta.mkdir(parents=True, exist_ok=True)
    escrever_csv_empresas(pasta / "COD_EMPRESAS.csv", [("ACME SAÚDE", "123"), ("BETA", "456")])

    criar_arquivo_pc(pasta / "PC123 OK.xlsx", TITULAR_OK, DEPENDENTE_OK, SUPLEMENTAR_OK)
    criar_arquivo_pc(
        pasta / "PC123 ABA FALTANDO.xlsx", TITULAR_OK, DEPENDENTE_OK, SUPLEMENTAR_OK,
        abas=("TITULAR", "DEPENDENTE", "OUTRA"),
    )
    criar_arquivo_pc(
        pasta / "PC123 DUPLICADO.xlsx", [TITULAR_OK[0], TITULAR_OK[0]], DEPENDENTE_OK, SUPLEMENTAR_OK
    )

    titulares_datas = [
        _com(TITULAR_OK[0], **{"DATA DE NASCIMENTO": "05/05/20",
                               "DATA DE NASCIMENTO Y2K": datetime(1920, 5, 5)}),
        _com(TITULAR_OK[1], **{"DATA DE NASCIMENTO": "01/01/90",
                               "DATA DE NASCIMENTO Y2K": "02/01/1990"}),
        _com(TITULAR_OK[1], **{"NUMERO DO CERTIFICADO": "003", "DATA DE NASCIMENTO": "xx",
                               "DATA DE NASCIMENTO Y2K": None}),
        _com(TITULAR_OK[1], **{"NUMERO DO CERTIFICADO": "004", "DATA DE NASCIMENTO": 45000,
                               "DATA DE NASCIMENTO Y2K": "31/02/71"}),
    ]
    dependentes_datas = [
        _com(DEPENDENTE_OK[0], **{"DATA DE NASCIMENTO": "2010-10-10"}),
        _com(DEPENDENTE_OK[1], **{"DATA DE NASCIMENTO": "20120201"}),
    ]
    criar_arquivo_pc(pasta / "PC456 DATAS.xlsx", titulares_datas, dependentes_datas, SUPLEMENTAR_OK)

    criar_arquivo_pc(
        pasta / "PC456 CONFLITO SUP.xlsx", TITULAR_OK, DEPENDENTE_OK,
        SUPLEMENTAR_OK + [_com(SUPLEMENTAR_OK[0], **{"CARTÃO": "OUTRO"})],
    )
    criar_arquivo_pc(pasta / "PC777 SEM EMPRESA.xlsx", TITULAR_OK, DEPENDENTE_OK, SUPLEMENTAR_OK)

    titulares_numericos = [
        _com(TITULAR_OK[0], **{"NUMERO DO CPF": 12345678901, "NUMERO DA SUBFATURA": 7}),
        _com(TITULAR_OK[1], **{"NUMERO DO CPF": 98765432100, "NUMERO DA SUBFATURA": 7.0,
                               "NUMERO DO CERTIFICADO": 2}),
    ]
    dependentes_numericos = [
        _com(DEPENDENTE_OK[0], **{"NUMERO DO CERTIFICADO": "2", "CODIGO DO DEPENDENTE": 1}),
    ]
    criar_arquivo_pc(
        pasta / "PC456 NUMERICO.xlsx", titulares_numericos, dependentes_numericos, SUPLEMENTAR_OK
    )

    criar_arquivo_pc(
        pasta / "PC123 LINHAS FORMATADAS.xlsx", TITULAR_OK, DEPENDENTE_OK, SUPLEMENTAR_OK,
        linhas_vazias_formatadas_no_fim=5,
    )

    # Achados da revisao final: <dimension> errado, NBSP e celula so com espaco.
    criar_arquivo_pc(pasta / "PC123 DIMENSAO.xlsx", TITULAR_OK, DEPENDENTE_OK, SUPLEMENTAR_OK)
    forcar_dimensao(pasta / "PC123 DIMENSAO.xlsx", 2, "A1")
    dependentes_espacos = [
        _com(DEPENDENTE_OK[0], **{"NUMERO DO CERTIFICADO": "001\xa0"}),
        DEPENDENTE_OK[1],
    ]
    p = pasta / "PC456 ESPACOS.xlsx"
    criar_arquivo_pc(p, TITULAR_OK, dependentes_espacos, SUPLEMENTAR_OK)
    wb = openpyxl.load_workbook(p)
    wb["DEPENDENTE"].append([" "])
    wb.save(p)

    # Revisao 2 (C2, C8, C11, C14): casos conferidos com o R.
    p = pasta / "PC123 TOTAL FORMULA.xlsx"
    criar_arquivo_pc(p, TITULAR_OK, DEPENDENTE_OK, SUPLEMENTAR_OK)
    wb = openpyxl.load_workbook(p)
    ws = wb["DEPENDENTE"]
    ws.cell(ws.max_row, 1).value = "=COUNTA(A2:A3)"  # gravado sem valor em cache
    wb.save(p)

    p = pasta / "PC123 TITULAR FORMULA VAZIA.xlsx"
    criar_arquivo_pc(p, TITULAR_OK, DEPENDENTE_OK, SUPLEMENTAR_OK)
    wb = openpyxl.load_workbook(p)
    ws = wb["TITULAR"]
    ws.cell(ws.max_row + 1, 1, "__E__")
    wb.save(p)
    trocar_celula(p, "__E__", FORMULA_TEXTO_CACHE_VAZIO)

    p = pasta / "PC123 TITULAR LINHA 3 VAZIA.xlsx"
    criar_arquivo_pc(p, TITULAR_OK, DEPENDENTE_OK, SUPLEMENTAR_OK)
    wb = openpyxl.load_workbook(p)
    wb["TITULAR"].cell(2, 1, "LINHA 2")
    wb["TITULAR"].insert_rows(3)
    wb.save(p)

    titulares_subf = [_com(TITULAR_OK[0], **{"NUMERO DA SUBFATURA": "0002"}),
                      _com(TITULAR_OK[1], **{"NUMERO DA SUBFATURA": "0001"})]
    # NOME DA MAE so com quebra de linha: NA nos dois (C8).
    dependentes_subf = [_com(d, **{"NUMERO DA SUBFATURA": "0009", "NOME DA MÃE DO DEPENDENTE": "\n"})
                        for d in DEPENDENTE_OK]
    criar_arquivo_pc(pasta / "PC456 SUBF.xlsx", titulares_subf, dependentes_subf, SUPLEMENTAR_OK)

    # P11: numeros gravados por exportador nao-Excel; coluna texto usa o <v> bruto.
    brutos = {
        "__CERT001__": '<v>001</v>', "__EXP__": '<v>1.2E3</v>', "__GRANDE__": '<v>12345678901234567</v>',
        "__DECIMAL__": '<v>0.3333333333333333</v>', "__DECIMAL2__": '<v>99999999.98999999</v>',
        "__MENOS0__": '<v>-0</v>', "__SERIAL60__": '<v>60</v>',
    }
    titulares_brutos = [
        _com(TITULAR_OK[0], **{"NUMERO DO CERTIFICADO": "__CERT001__", "NUMERO DA SUBFATURA": "__EXP__",
                               "NUMERO DO CPF": "__GRANDE__", "REGIÃO": "__DECIMAL__",
                               "SEXO DO SEGURADO": "__ISO__", "DATA DE NASCIMENTO": "__ISO__"}),
        _com(TITULAR_OK[1], **{"NUMERO DO CPF": "__DECIMAL2__", "REGIÃO": "__MENOS0__",
                               "SEXO DO SEGURADO": "__SERIAL60__"}),
    ]
    p = pasta / "PC456 TEXTO BRUTO.xlsx"
    criar_arquivo_pc(p, titulares_brutos, DEPENDENTE_OK, SUPLEMENTAR_OK)
    for marcador, v in brutos.items():
        trocar_celula(p, marcador, lambda ref, v=v: f'<c r="{ref}">{v}</c>')
    trocar_celula(p, "__ISO__", lambda ref: f'<c r="{ref}" t="d"><v>2023-03-15T00:00:00</v></c>')
    return pasta
