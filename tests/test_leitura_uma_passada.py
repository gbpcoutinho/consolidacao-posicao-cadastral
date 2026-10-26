"""Leitura em uma passada (fase 2) = leitura antiga (openpyxl + varredura) em
XML incomum: linhas e celulas fora de ordem ou sem r, texto rico, tipos raros."""
import zipfile

import openpyxl
import pytest

from consolidacao import leitura
from tests import leitura_antiga

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

LINHAS_XML = {
    "sem_r": '<row><c t="inlineStr"><is><t>A</t></is></c><c><v>7</v></c></row>'
             '<row r="4"><c r="C4"><f>1+1</f></c><c><v/></c></row>',
    "linha_repetida_e_fora_de_ordem": '<row r="3"><c r="A3" t="inlineStr"><is><t>x</t></is></c></row>'
                                      '<row r="2"><c r="A2"><v>1</v></c></row>'
                                      '<row r="3"><c r="B3"><v>2</v></c></row>'
                                      '<row r="5"/>',
    "celulas_fora_de_ordem": '<row r="1"><c r="C1"><v>3</v></c><c r="A1"><v>1</v></c></row>'
                             '<row r="2"><c r="B2"><v>5</v></c><c r="B2"><v>6</v></c></row>',
    "tipos": '<row r="1"><c r="A1" t="b"><v>1</v></c><c r="B1" t="e"><v>#N/A</v></c>'
             '<c r="C1" t="str"><f>A1</f><v>  ab  </v></c><c r="D1" t="d"><v>2023-03-15T00:00:00</v></c>'
             '<c r="E1"><v>1.5E3</v></c><c r="F1"><v>0012</v></c><c r="G1" t="str"><v> </v></c></row>',
    "texto_rico": '<row r="1"><c r="A1" t="inlineStr"><is><r><t>ab</t></r><r><t xml:space="preserve"> cd</t></r>'
                  '</is></c><c r="B1" t="inlineStr"><is><t xml:space="preserve">  e  </t></is></c>'
                  '<c r="C1" t="inlineStr"><is><t/></is></c><c r="D1" t="inlineStr"/></row>',
    "estilo_data": '<row r="1"><c r="A1" s="1"><v>45000</v></c><c r="B1" s="1"><v>9999999999</v></c>'
                   '<c r="C1" s="1"><v>0.5</v></c></row>',
}


def _planilha(tmp_path, nome, linhas_xml):
    base = tmp_path / f"{nome}_base.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "A"
    ws["A1"] = "x"
    ws["B1"] = 1
    ws["B1"].number_format = "dd/mm/yyyy"  # cria o estilo s="1" com formato de data
    wb.save(base)
    xml = f'<worksheet xmlns="{NS}"><sheetData>{linhas_xml}</sheetData></worksheet>'
    destino = tmp_path / f"{nome}.xlsx"
    with zipfile.ZipFile(base) as origem, zipfile.ZipFile(destino, "w") as saida:
        for item in origem.infolist():
            dados = origem.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                dados = xml.encode()
            saida.writestr(item, dados)
    return destino


@pytest.mark.parametrize("nome", sorted(LINHAS_XML))
def test_mesmas_linhas_que_a_leitura_antiga(tmp_path, nome):
    caminho = _planilha(tmp_path, nome, LINHAS_XML[nome])
    antigo = leitura_antiga.normalizar(leitura_antiga.linhas_brutas(caminho, "A"))
    wb = leitura.abrir_planilha(caminho)
    try:
        linhas, varridas = leitura._Leitor(wb, "A").ler()
    finally:
        wb.close()
    leitura._aplicar_varredura(linhas, varridas)
    assert leitura_antiga.normalizar(linhas) == antigo


@pytest.mark.parametrize("r", ["2.5", "2.0", "x"])
def test_linha_r_invalido_mesmo_erro(tmp_path, r):
    caminho = _planilha(tmp_path, "r_invalido", f'<row r="{r}"><c r="A2"><v>1</v></c></row>')
    with pytest.raises(ValueError) as antigo:
        leitura_antiga.linhas_brutas(caminho, "A")
    wb = leitura.abrir_planilha(caminho)
    try:
        with pytest.raises(ValueError) as novo:
            leitura._Leitor(wb, "A").ler()
    finally:
        wb.close()
    assert str(novo.value) == str(antigo.value)
