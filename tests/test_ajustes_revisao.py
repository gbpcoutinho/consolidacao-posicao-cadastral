"""Achados do debate dos 3 revisores (2026-10-22)."""
from datetime import date

import openpyxl

from consolidacao.leitura import ler_aba
from consolidacao.datas import validar_colunas_data


def _planilha_formatada(caminho, cabecalho, valores, formato):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "A"
    ws.append(cabecalho)
    ws.append(valores)
    for celula in ws[2]:
        celula.number_format = formato
    wb.save(caminho)


def test_a1_numero_grande_em_celula_de_data_nao_some(tmp_path):
    p = tmp_path / "x.xlsx"
    _planilha_formatada(p, ["NUMERO DO CPF", "DATA DE NASCIMENTO"], [12345678901, 19900512], "dd/mm/yyyy")
    df = ler_aba(p, "A")
    assert df["NUMERO DO CPF"].tolist() == ["12345678901"]
    _, obs, _ = validar_colunas_data(df, "C")
    assert len(obs) == 1 and "19900512" in obs[0]


def test_a1_data_normal_continua_data(tmp_path):
    p = tmp_path / "x.xlsx"
    _planilha_formatada(p, ["DATA DE NASCIMENTO"], [45000], "dd/mm/yyyy")
    df, obs, _ = validar_colunas_data(ler_aba(p, "A"), "C")
    assert df["DATA DE NASCIMENTO"].tolist() == [date(2023, 3, 15)] and obs == []


def _consolidado_minimo(tmp_path):
    from consolidacao.processamento import processar_arquivos
    from tests.fabrica import DEPENDENTE_OK, SUPLEMENTAR_OK, TITULAR_OK, criar_arquivo_pc

    a = tmp_path / "PC123 x.xlsx"
    criar_arquivo_pc(a, TITULAR_OK, DEPENDENTE_OK, SUPLEMENTAR_OK)
    return processar_arquivos([a], {"123": "ACME"})


def test_x1_salvar_cria_pasta_e_nao_deixa_temporario(tmp_path):
    from consolidacao.exportacao import escrever_excel

    destino = tmp_path / "nova" / "saida.xlsx"
    escrever_excel(_consolidado_minimo(tmp_path), destino)
    assert destino.exists()
    assert [p.name for p in destino.parent.iterdir()] == ["saida.xlsx"]


def test_x1_falha_ao_salvar_vira_erro_claro_sem_temporario(tmp_path):
    import pytest

    from consolidacao.erros import ErroValidacao
    from consolidacao.exportacao import escrever_excel

    c = _consolidado_minimo(tmp_path)
    destino = tmp_path / "saida.xlsx"
    destino.mkdir()  # destino ocupado: os.replace falha, como com o arquivo aberto no Excel
    with pytest.raises(ErroValidacao, match="Nao foi possivel salvar o arquivo de saida"):
        escrever_excel(c, destino)
    assert sorted(p.name for p in tmp_path.iterdir()) == ["PC123 x.xlsx", "saida.xlsx"]


def test_x1_cli_retorna_1_quando_nao_salva(tmp_path, capsys):
    from tests.fabrica import DEPENDENTE_OK, SUPLEMENTAR_OK, TITULAR_OK, criar_arquivo_pc, escrever_csv_empresas
    from tratamento import main

    criar_arquivo_pc(tmp_path / "PC123 x.xlsx", TITULAR_OK, DEPENDENTE_OK, SUPLEMENTAR_OK)
    escrever_csv_empresas(tmp_path / "e.csv", [("ACME", "123")])
    (tmp_path / "ocupado.xlsx").mkdir()
    codigo = main(["--entrada", str(tmp_path), "--empresas", str(tmp_path / "e.csv"),
                   "--saida", str(tmp_path / "ocupado.xlsx")])
    assert codigo == 1
    assert "Erro: Nao foi possivel salvar o arquivo de saida" in capsys.readouterr().out


def _dois_arquivos(tmp_path):
    from tests.fabrica import DEPENDENTE_OK, SUPLEMENTAR_OK, TITULAR_OK, criar_arquivo_pc

    arquivos = [tmp_path / "PC123 a.xlsx", tmp_path / "PC123 b.xlsx"]
    for a in arquivos:
        criar_arquivo_pc(a, TITULAR_OK, DEPENDENTE_OK, SUPLEMENTAR_OK)
    return arquivos


def test_x2_progresso_recebe_indice_total_e_nome(tmp_path):
    from consolidacao.processamento import processar_arquivos

    vistos = []
    processar_arquivos(_dois_arquivos(tmp_path), {"123": "A"},
                       progresso=lambda i, total, nome: vistos.append((i, total, nome)))
    assert vistos == [(1, 2, "PC123 a.xlsx"), (2, 2, "PC123 b.xlsx")]


def test_x2_cancelar_interrompe_antes_do_proximo_arquivo(tmp_path):
    import pytest

    from consolidacao.erros import ProcessamentoCancelado
    from consolidacao.processamento import processar_arquivos

    vistos = []
    with pytest.raises(ProcessamentoCancelado):
        processar_arquivos(_dois_arquivos(tmp_path), {"123": "A"},
                           progresso=lambda i, total, nome: vistos.append(i),
                           cancelar=lambda: len(vistos) >= 1)
    assert vistos == [1]


def test_c1_executar_processa_grava_e_resume(tmp_path):
    from consolidacao.pipeline import executar, resumir

    saida = tmp_path / "out" / "r.xlsx"
    c = executar(_dois_arquivos(tmp_path), {"123": "A"}, saida)
    assert saida.exists()
    assert resumir(c) == {
        "arquivos_ok": 2, "arquivos_erro": 0, "registros": 8, "pendencias": 2,
        "datas_invalidas": 0, "divergencias_datas": 0, "erros": [],
    }


def test_c1_cli_usa_data_injetada_no_nome_padrao(tmp_path):
    from datetime import date

    from tests.fabrica import escrever_csv_empresas
    from tratamento import main

    _dois_arquivos(tmp_path)
    escrever_csv_empresas(tmp_path / "e.csv", [("ACME", "123")])
    assert main(["--entrada", str(tmp_path), "--empresas", str(tmp_path / "e.csv")],
                hoje=date(2027, 1, 5)) == 0
    assert (tmp_path / "PC_CONSOLIDADO_1_2027.xlsx").exists()


def _excel_de(tmp_path, df, nome="Base Tratada"):
    import pandas as pd

    from consolidacao.exportacao import escrever_excel
    from consolidacao.processamento import Consolidado

    vazio = pd.DataFrame({"X": pd.Series([], dtype=object)})
    c = Consolidado(consolidado=df, pendencias=vazio, log=vazio)
    destino = tmp_path / "o.xlsx"
    escrever_excel(c, destino)
    return openpyxl.load_workbook(destino)[nome]


def test_b2_caractere_de_controle_nao_quebra_exportacao(tmp_path):
    import pandas as pd

    ws = _excel_de(tmp_path, pd.DataFrame({"NOME": ["a\x0bb"]}, dtype=object))
    assert ws["A2"].value == "ab"


def test_b5_largura_automatica_tem_teto(tmp_path):
    import pandas as pd

    ws = _excel_de(tmp_path, pd.DataFrame({"TEXTO": ["x" * 300]}, dtype=object))
    assert ws.column_dimensions["A"].width == 100


def test_a4_coluna_de_data_vazia_recebe_formato_de_data(tmp_path):
    import pandas as pd

    df = pd.DataFrame({"NOME": ["a"], "DATA DE REATIVACAO": [None]}, dtype=object)
    ws = _excel_de(tmp_path, df)
    assert ws["B2"].number_format == "dd/mm/yyyy"
    assert ws["B2"].alignment.horizontal == "center"


def test_a3_observacao_com_quebra_de_linha_agrupa():
    from consolidacao.datas import resumir_observacoes_datas

    obs = ["C | coluna 'D' | registro 1 | valor invalido 'a\nb' (convertido para vazio)"]
    assert resumir_observacoes_datas(obs) == (
        "C | coluna 'D': 1 valor(es); amostras: registro 1 | valor invalido 'a\nb' (convertido para vazio)"
    )


def test_b3_cada_arquivo_e_aberto_uma_vez(tmp_path, monkeypatch):
    import consolidacao.leitura as leitura
    import consolidacao.processamento as processamento

    chamadas = []
    original = openpyxl.load_workbook

    def contar(*args, **kwargs):
        chamadas.append(args[0])
        return original(*args, **kwargs)

    monkeypatch.setattr(leitura.openpyxl, "load_workbook", contar)
    monkeypatch.setattr(processamento.openpyxl, "load_workbook", contar)
    processamento.processar_arquivos(_dois_arquivos(tmp_path)[:1], {"123": "A"})
    assert len(chamadas) == 1


def test_c6_cli_retorna_0_mesmo_com_todos_os_arquivos_em_erro(tmp_path):
    from tests.fabrica import escrever_csv_empresas
    from tratamento import main

    _dois_arquivos(tmp_path)
    escrever_csv_empresas(tmp_path / "e.csv", [("OUTRA", "999")])
    # Como no R: o processamento termina e os erros ficam no Log.
    assert main(["--entrada", str(tmp_path), "--empresas", str(tmp_path / "e.csv")]) == 0


def test_a2_duracao_em_coluna_de_texto_vira_numero(tmp_path):
    p = tmp_path / "x.xlsx"
    _planilha_formatada(p, ["HORAS"], [1.5], "[h]:mm:ss")
    assert ler_aba(p, "A")["HORAS"].tolist() == ["1.5"]
