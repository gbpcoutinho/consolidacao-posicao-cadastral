from datetime import datetime

import openpyxl

from consolidacao.exportacao import escrever_excel
from consolidacao.processamento import processar_arquivos
from tests.fabrica import DEPENDENTE_OK, SUPLEMENTAR_OK, TITULAR_OK, criar_arquivo_pc


def test_excel_formatado(tmp_path):
    a = tmp_path / "PC123 x.xlsx"
    criar_arquivo_pc(a, TITULAR_OK, DEPENDENTE_OK, SUPLEMENTAR_OK)
    destino = tmp_path / "out.xlsx"
    escrever_excel(processar_arquivos([a], {"123": "ACME"}), destino)
    wb = openpyxl.load_workbook(destino)
    assert wb.sheetnames == ["Base Tratada", "Pendencias", "Log"]
    ws = wb["Base Tratada"]
    assert "TabelaBaseTratada" in ws.tables
    assert ws.tables["TabelaBaseTratada"].tableStyleInfo.name == "TableStyleMedium2"
    assert ws.freeze_panes == "A2" and ws.row_dimensions[1].height == 32
    assert ws["A1"].fill.fgColor.rgb.endswith("4472C4") and ws["A1"].font.bold
    assert ws.sheet_view.showGridLines is False
    col = {c.value: c.column_letter for c in ws[1]}
    assert ws.column_dimensions[col["EMPRESA"]].width == 28
    assert ws.column_dimensions[col["DATA DE NASCIMENTO"]].width == 15
    celula = ws[f'{col["DATA DE NASCIMENTO"]}2']
    assert celula.number_format == "dd/mm/yyyy" and isinstance(celula.value, datetime)
    assert ws[f'{col["CODIGO DO DEPENDENTE"]}2'].alignment.horizontal == "right"
    assert ws[f'{col["PLANO"]}2'].alignment.horizontal == "center"
    log = wb["Log"]
    assert [c.value for c in log[1]][:3] == ["ARQUIVO", "REFERENCIA", "CODIGO_EMPRESA"]


def test_largura_de_data_cabe_o_cabecalho_em_duas_linhas():
    from consolidacao.exportacao import largura_coluna_data

    # "DATA DE INICIO" / "DE VIGENCIA Y2K": a linha mais longa tem 15 caracteres.
    assert 20 <= largura_coluna_data("DATA DE INICIO DE VIGENCIA Y2K") <= 22
    # "DATA DE" / "NASCIMENTO": 10 caracteres.
    assert 13 <= largura_coluna_data("DATA DE NASCIMENTO") <= 15
    # Cabecalhos curtos mantem o minimo de 13 do R.
    assert largura_coluna_data("DATA") == 13


def test_excel_pendencias_vazias_sem_tabela(tmp_path):
    a = tmp_path / "PC123 x.xlsx"
    criar_arquivo_pc(a, TITULAR_OK, DEPENDENTE_OK[:1], SUPLEMENTAR_OK)
    destino = tmp_path / "out.xlsx"
    escrever_excel(processar_arquivos([a], {"123": "ACME"}), destino)
    ws = openpyxl.load_workbook(destino)["Pendencias"]
    assert not ws.tables and ws.auto_filter.ref == "A1:I1"
