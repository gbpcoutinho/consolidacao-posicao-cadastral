"""O gravador da saida (otimizado na fase 2) produz o mesmo arquivo que o
gravador anterior: valores, estilos, larguras, tabela, filtro e congelamento."""
from pathlib import Path

import openpyxl
import pytest

from consolidacao.empresas import carregar_mapa_empresas
from consolidacao.exportacao import escrever_excel
from consolidacao.processamento import processar_arquivos
from tests import exportacao_antiga
from tests.fabrica import DEPENDENTE_OK, SUPLEMENTAR_OK, TITULAR_OK, criar_arquivo_pc

EXEMPLOS = Path(__file__).resolve().parents[1] / "exemplos"


def _estilo(c):
    return (
        c.number_format,
        (c.alignment.horizontal, c.alignment.vertical, bool(c.alignment.wrap_text)),
        (bool(c.font.bold), c.font.color.rgb if c.font.color is not None else None),
        (c.fill.fill_type, c.fill.fgColor.rgb if c.fill.fill_type else None),
        c.border.bottom.style if c.border.bottom is not None else None,
    )


def _descrever(caminho):
    wb = openpyxl.load_workbook(caminho)
    folhas = {}
    for ws in wb.worksheets:
        tabelas = {
            t.displayName: (t.ref, t.tableStyleInfo.name, t.tableStyleInfo.showRowStripes,
                            [col.name for col in t.tableColumns],
                            t.autoFilter.ref if t.autoFilter is not None else None)
            for t in ws.tables.values()
        }
        larguras = {
            openpyxl.utils.get_column_letter(j): ws.column_dimensions[
                openpyxl.utils.get_column_letter(j)
            ].width
            for j in range(1, ws.max_column + 1)
        }
        celulas = [
            [(c.value, _estilo(c)) for c in linha]
            for linha in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column)
        ]
        folhas[ws.title] = {
            "dimensoes": (ws.max_row, ws.max_column),
            "congelado": ws.freeze_panes,
            "grade": ws.sheet_view.showGridLines,
            "filtro": ws.auto_filter.ref,
            "tabelas": tabelas,
            "larguras": larguras,
            "altura1": ws.row_dimensions[1].height,
            "celulas": celulas,
        }
    return wb.sheetnames, folhas


def _comparar(consolidado, tmp_path):
    antigo, novo = tmp_path / "antigo.xlsx", tmp_path / "novo.xlsx"
    exportacao_antiga.escrever_excel(consolidado, antigo)
    escrever_excel(consolidado, novo)
    nomes_a, folhas_a = _descrever(antigo)
    nomes_b, folhas_b = _descrever(novo)
    assert nomes_a == nomes_b
    for nome in nomes_a:
        a, b = folhas_a[nome], folhas_b[nome]
        for chave in a:
            if chave == "celulas":
                for i, (la, lb) in enumerate(zip(a[chave], b[chave]), start=1):
                    assert la == lb, f"[{nome}] linha {i}"
            assert a[chave] == b[chave], f"[{nome}] {chave}"


def test_exemplos_iguais_ao_gravador_anterior(tmp_path):
    mapa = carregar_mapa_empresas(EXEMPLOS / "COD_EMPRESAS_EXEMPLO.csv")
    arquivos = sorted((EXEMPLOS / "entrada").glob("PC9000*.xlsx"))
    assert len(arquivos) == 3
    _comparar(processar_arquivos(arquivos, mapa), tmp_path)


@pytest.mark.parametrize("dependentes", [DEPENDENTE_OK, DEPENDENTE_OK[:1]])
def test_sinteticos_iguais_ao_gravador_anterior(tmp_path, dependentes):
    # DEPENDENTE_OK[:1]: Pendencias vazia (sem tabela, so filtro).
    a = tmp_path / "PC123 x.xlsx"
    criar_arquivo_pc(a, TITULAR_OK, dependentes, SUPLEMENTAR_OK)
    b = tmp_path / "PC999 y.xlsx"  # arquivo com erro: Log com colunas vazias
    criar_arquivo_pc(b, TITULAR_OK, dependentes, SUPLEMENTAR_OK)
    _comparar(processar_arquivos([a, b], {"123": "ACME"}), tmp_path)


def test_base_vazia_igual_ao_gravador_anterior(tmp_path):
    b = tmp_path / "PC999 y.xlsx"
    criar_arquivo_pc(b, TITULAR_OK, DEPENDENTE_OK, SUPLEMENTAR_OK)
    _comparar(processar_arquivos([b], {"1": "A"}), tmp_path)


def test_openpyxl_usa_lxml_para_gravar():
    # lxml (requirements.txt) deixa a serializacao do XML ~1/3 mais rapida.
    assert openpyxl.LXML is True
